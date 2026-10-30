"""Cash Flow Intelligence: CFO quality, CapEx intensity, FCF, distress, deleveraging, capital allocation."""
from __future__ import annotations

import logging

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.analytics.ratios import PATTERNS, build_panel, cagr
from src.etl.clean import OUT

logger = logging.getLogger(__name__)
WINDOW = 5
COLS = ["company_id", "sector", "cfo_quality_score", "cfo_quality_label", "capex_intensity_pct", "capex_label",
        "fcf_cagr_5yr", "fcf_conversion_pct", "distress_flag", "deleveraging_flag", "capital_allocation_label"]


def cfo_label(x: float) -> str:
    """High Quality (>1.0), Moderate (0.5-1.0), Accrual Risk (<0.5)."""
    if pd.isna(x):
        return "n/a"
    return "High Quality" if x > 1.0 else ("Moderate" if x >= 0.5 else "Accrual Risk")


def capex_label(x: float) -> str:
    """Asset Light (<3%), Moderate (3-8%), Capital Intensive (>8%)."""
    if pd.isna(x):
        return "n/a"
    return "Asset Light" if x < 3 else ("Moderate" if x <= 8 else "Capital Intensive")


def company_kpis(g: pd.DataFrame) -> dict:
    """Compute all cash-flow KPIs for one company's panel."""
    g = g.sort_values("fy")
    cf = g[g.operating_activity.notna()]
    r = {"company_id": g.company_id.iloc[0], "sector": g.broad_sector.iloc[0], "notes": []}
    if cf.empty:
        r.update(cfo_quality_score=np.nan, cfo_quality_label="n/a", capex_intensity_pct=np.nan, capex_label="n/a",
                 fcf_cagr_5yr=np.nan, fcf_conversion_pct=np.nan, distress_flag=False, deleveraging_flag=False,
                 capital_allocation_label="Insufficient data", latest_year=None, cfo=np.nan, cff=np.nan,
                 net_profit=np.nan, fcf_cagr_flag="INSUFFICIENT")
        r["notes"].append("no cash-flow statement in source")
        r["notes"] = "; ".join(r["notes"])
        return r
    w = cf.tail(WINDOW)
    ratios = (w.operating_activity / w.net_profit.where(w.net_profit > 0)).dropna()
    r["cfo_quality_score"] = ratios.mean() if len(ratios) else np.nan
    if len(ratios) < len(w):
        r["notes"].append(f"CFO/PAT undefined in {len(w)-len(ratios)} loss yr(s)")
    r["cfo_quality_label"] = cfo_label(r["cfo_quality_score"])
    ws = w[w.sales > 0]
    r["capex_intensity_pct"] = ws.investing_activity.abs().sum() / ws.sales.sum() * 100 if len(ws) else np.nan
    r["capex_label"] = capex_label(r["capex_intensity_pct"])
    r["fcf_cagr_5yr"], r["fcf_cagr_flag"] = cagr(cf.fcf, cf.fy, 5)
    we = w[w.operating_profit > 0]
    r["fcf_conversion_pct"] = we.fcf.sum() / we.operating_profit.sum() * 100 if len(we) else np.nan
    last = cf.iloc[-1]
    r["latest_year"], r["cfo"], r["cff"], r["net_profit"] = last.year, last.operating_activity, last.financing_activity, last.net_profit
    r["distress_flag"] = bool(last.operating_activity < 0 and last.financing_activity > 0)
    prev = g[g.fy < last.fy].tail(1)
    delev = False
    if len(prev) and not pd.isna(prev.borrowings.iloc[0]) and not pd.isna(last.borrowings):
        delev = bool(last.financing_activity < 0 and last.borrowings < prev.borrowings.iloc[0])
    r["deleveraging_flag"] = delev
    r["capital_allocation_label"] = last.pattern_label or "Insufficient data"
    if bool(last.is_financial):
        r["notes"].append("financial: CFO/CFI/CFF include loan-book & deposit flows - interpret with care")
    r["notes"] = "; ".join(r["notes"])
    return r


def pattern_changes(p: pd.DataFrame) -> pd.DataFrame:
    """Every year-over-year change in capital-allocation label."""
    rows = []
    for cid, g in p[p.pattern_label.notna()].groupby("company_id"):
        g = g.sort_values("fy")
        for (_, a), (_, b) in zip(g.iloc[:-1].iterrows(), g.iloc[1:].iterrows()):
            if a.pattern_label != b.pattern_label:
                rows.append(dict(company_id=cid, sector=a.broad_sector, from_year=a.year, to_year=b.year,
                                 from_pattern=a.pattern_label, to_pattern=b.pattern_label))
    return pd.DataFrame(rows)


def verify_capital_allocation(p: pd.DataFrame, companies: list[str]) -> str:
    """Completeness check of capital_allocation.csv (companies x years)."""
    ca = pd.read_csv(OUT / "capital_allocation.csv")
    lines = [f"capital_allocation.csv rows: {len(ca)}", f"companies covered: {ca.company_id.nunique()} / {len(companies)}",
             f"rows with a pattern label: {ca.pattern_label.notna().sum()}  | unlabelled: {ca.pattern_label.isna().sum()}"]
    miss = ca[ca.pattern_label.isna()].groupby("company_id").year.agg(["count", "min", "max"])
    lines.append("companies with unlabelled years (no cash-flow data in source):")
    lines += [f"  {i}: {r['count']} yrs ({r['min']}..{r['max']})" for i, r in miss.iterrows()]
    lines.append("companies with ZERO labelled years: " + ", ".join(
        sorted(set(companies) - set(ca[ca.pattern_label.notna()].company_id))) or "none")
    return "\n".join(lines)


def style_sheet(ws, widths: dict[str, int] | None = None) -> None:
    """Header style, freeze panes, column widths."""
    hdr = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill, c.alignment = hdr, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
    for i, col in enumerate(ws.columns, 1):
        mx = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(12, mx + 2))
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 32


def run() -> pd.DataFrame:
    """Write cashflow_intelligence.xlsx, distress_alerts.csv, pattern_changes.csv, verification text."""
    p = build_panel()
    cos = pd.read_excel(OUT.parent / "data/raw/companies.xlsx", header=1).id.str.strip().str.upper().tolist()
    k = pd.DataFrame([company_kpis(p[p.company_id == c]) if (p.company_id == c).any() else
                      {"company_id": c, "notes": "no P&L"} for c in cos])
    main = k[COLS + ["fcf_cagr_flag", "notes"]].rename(columns={"notes": "data_notes"}).copy()
    for c, d in (("cfo_quality_score", 2), ("capex_intensity_pct", 2), ("fcf_cagr_5yr", 2), ("fcf_conversion_pct", 1)):
        main[c] = main[c].astype(float).round(d)

    dist = k[k.distress_flag.fillna(False)][["company_id", "sector", "latest_year", "cfo", "cff", "net_profit"]].copy()
    dist.columns = ["company_id", "sector", "year", "cfo_value", "cff_value", "latest_net_profit"]
    dist["context"] = np.where(dist.sector.eq("Financials"),
                               "Financial company: negative CFO / positive CFF is structural (loan growth funded by deposits/borrowings), not necessarily distress",
                               "Operating cash burn funded by financing - review")
    dist = dist.sort_values(["context", "cfo_value"])
    dist.to_csv(OUT / "distress_alerts.csv", index=False)

    changes = pattern_changes(p)
    changes.to_csv(OUT / "pattern_changes.csv", index=False)
    (OUT / "capital_allocation_verification.txt").write_text(verify_capital_allocation(p, cos))

    dist_summary = main.capital_allocation_label.value_counts().rename_axis("pattern").reset_index(name="companies")
    order = list(PATTERNS.values()) + ["Insufficient data"]
    dist_summary["o"] = dist_summary.pattern.map({v: i for i, v in enumerate(order)})
    dist_summary = dist_summary.sort_values("o").drop(columns="o")
    legend = pd.DataFrame({"Item": ["CFO quality", "CapEx intensity", "FCF CAGR", "FCF conversion", "Distress", "Deleveraging",
                                    "Capital allocation", "Window"],
                           "Definition": ["Mean of CFO/PAT over last 5 yrs with PAT>0. >1.0 High Quality, 0.5-1.0 Moderate, <0.5 Accrual Risk",
                                          "5-yr sum |investing activity| / 5-yr sum sales x100. <3 Asset Light, 3-8 Moderate, >8 Capital Intensive",
                                          "5-yr CAGR of FCF (CFO+CFI); blank when base<=0 / insufficient history (see fcf_cagr_flag)",
                                          "5-yr sum FCF / 5-yr sum operating profit x100 (years with op. profit > 0)",
                                          "Latest cash-flow year: CFO<0 AND CFF>0",
                                          "Latest cash-flow year: CFF<0 AND borrowings lower than prior year",
                                          "Sign pattern of CFO/CFI/CFF in latest cash-flow year (8 classes)",
                                          "'Latest year' = last year with a cash-flow statement for that company"]})
    flags = k[["company_id", "fcf_cagr_flag"]].fillna("")
    with pd.ExcelWriter(OUT / "cashflow_intelligence.xlsx", engine="openpyxl") as xw:
        main.to_excel(xw, sheet_name="Cashflow Intelligence", index=False)
        dist_summary.to_excel(xw, sheet_name="Pattern Distribution", index=False)
        changes.to_excel(xw, sheet_name="Pattern Changes", index=False)
        dist.to_excel(xw, sheet_name="Distress Alerts", index=False)
        flags.to_excel(xw, sheet_name="FCF CAGR Flags", index=False)
        legend.to_excel(xw, sheet_name="Legend", index=False)
        for ws in xw.book.worksheets:
            style_sheet(ws)
        ws = xw.book["Cashflow Intelligence"]
        red, grn = PatternFill("solid", fgColor="F8CBAD"), PatternFill("solid", fgColor="C6E0B4")
        hdr = {c.value: c.column for c in ws[1]}
        for row in ws.iter_rows(min_row=2):
            for name, bad, good in (("cfo_quality_label", "Accrual Risk", "High Quality"),
                                    ("distress_flag", True, None), ("deleveraging_flag", None, True)):
                cell = row[hdr[name] - 1]
                if bad is not None and cell.value == bad:
                    cell.fill = red
                if good is not None and cell.value == good:
                    cell.fill = grn
    return main


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    m = run()
    print(m.shape, m.company_id.nunique())
    print(m.cfo_quality_label.value_counts().to_dict(), m.capex_label.value_counts().to_dict())
    print(m.capital_allocation_label.value_counts().to_dict(), "distress", int(m.distress_flag.sum()), "delev", int(m.deleveraging_flag.sum()))
    print(open(OUT / "capital_allocation_verification.txt").read())
    print(m.head(5).to_string())
