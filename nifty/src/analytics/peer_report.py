"""peer_comparison.xlsx — one sheet per peer group, 20 metrics + 20 percentile ranks, colour-coded."""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .peer import ALL_METRICS, compute_percentiles

GREEN = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
YELLOW = PatternFill("solid", start_color="FFEB9C", end_color="FFEB9C")
RED = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
GOLD = PatternFill("solid", start_color="FFD966", end_color="FFD966")
NAVY = PatternFill("solid", start_color="0B2239", end_color="0B2239")
TEAL = PatternFill("solid", start_color="1F6F5C", end_color="1F6F5C")
GREY = PatternFill("solid", start_color="E7EBEF", end_color="E7EBEF")
_THIN = Side(style="thin", color="D0D7DE")
_FMT = {"sales": "#,##0", "net_profit": "#,##0", "fcf": "#,##0", "pe": "0.0", "pb": "0.00",
        "de": "0.00", "cfo_pat": "0.00", "asset_turnover": "0.00", "div_yield": "0.00"}


def percentile_fill(p: float | None):
    """>= 75th green, <= 25th red, otherwise yellow (boundaries: 0.75 -> green, 0.25 -> red)."""
    if p is None or pd.isna(p):
        return None
    if p >= 0.75:
        return GREEN
    if p <= 0.25:
        return RED
    return YELLOW


def write_peer_workbook(scored: pd.DataFrame, peer_groups: pd.DataFrame, path: str | Path) -> Path:
    pct = compute_percentiles(scored, peer_groups, ALL_METRICS)
    pct_wide = pct.pivot(index=["peer_group_name", "company_id"], columns="metric", values="percentile_rank")
    snap = scored.set_index("company_id")
    n_metrics = len(ALL_METRICS)

    wb = Workbook()
    wb.remove(wb.active)
    for group, members in peer_groups.groupby("peer_group_name", sort=True):
        ws = wb.create_sheet(group[:31])
        bench = set(members.loc[members["is_benchmark"], "company_id"])
        ids = [c for c in members["company_id"] if c in snap.index]
        ids.sort(key=lambda c: (c not in bench, -(snap.loc[c, "composite_quality_score"] or 0)))

        ws["A1"] = f"{group} — peer comparison (latest FY)"
        ws["A1"].font = Font(bold=True, size=14, color="0B2239")
        ws["A2"] = ("Percentile rank = PERCENT_RANK within the peer group (higher = better; D/E, P/E and P/B inverted so lower "
                    "values rank higher).  Green >= 75th, yellow 25th-75th, red <= 25th.  Gold row = benchmark company.  P/E, P/B, dividend yield are SIMULATED (market_cap dataset).")
        ws["A2"].font = Font(italic=True, size=9, color="555555")

        top, hdr = 4, 5
        ws.cell(row=top, column=3, value="METRIC VALUES").font = Font(bold=True, color="FFFFFF")
        ws.cell(row=top, column=3 + n_metrics, value="PERCENTILE RANK (within peer group)").font = Font(bold=True, color="FFFFFF")
        for j in range(1, 3 + n_metrics):
            ws.cell(row=top, column=j).fill = NAVY
        for j in range(3 + n_metrics, 3 + 2 * n_metrics):
            ws.cell(row=top, column=j).fill = TEAL

        heads = ["company_id", "company_name"] + [lab for _, lab, _ in ALL_METRICS] + \
                [f"{lab} — pctile" for _, lab, _ in ALL_METRICS]
        for j, h in enumerate(heads, 1):
            c = ws.cell(row=hdr, column=j, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = NAVY if j <= 2 + n_metrics else TEAL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[hdr].height = 44

        for i, cid in enumerate(ids, 1):
            r = hdr + i
            is_b = cid in bench
            ws.cell(row=r, column=1, value=cid)
            ws.cell(row=r, column=2, value=str(snap.loc[cid, "company_name"]) + ("  ★ benchmark" if is_b else ""))
            for j, (key, _, _) in enumerate(ALL_METRICS, 3):
                v = snap.loc[cid, key]
                shown = "Debt Free" if key == "icr" and isinstance(v, (float, np.floating)) and np.isinf(v) else v
                if isinstance(shown, (float, np.floating)) and np.isnan(shown):
                    shown = None
                c = ws.cell(row=r, column=j, value=shown)
                c.number_format = _FMT.get(key, "0.0")
            for j, (key, _, _) in enumerate(ALL_METRICS, 3 + n_metrics):
                p = pct_wide.loc[(group, cid), key]
                c = ws.cell(row=r, column=j, value=None if pd.isna(p) else float(p))
                c.number_format = "0%"
                c.alignment = Alignment(horizontal="center")
                f = percentile_fill(p)
                if f:
                    c.fill = f
            if is_b:
                for j in range(1, 3 + n_metrics):
                    ws.cell(row=r, column=j).fill = GOLD
                    ws.cell(row=r, column=j).font = Font(bold=True)
            for j in range(1, 3 + 2 * n_metrics):
                ws.cell(row=r, column=j).border = Border(bottom=_THIN)

        # summary row: peer-group median of each metric (Debt Free ICR counted as +inf)
        s = hdr + len(ids) + 1
        ws.cell(row=s, column=1, value="Peer group median")
        ws.cell(row=s, column=2, value=f"{group} (n={len(ids)})")
        sub = snap.loc[ids]
        for j, (key, _, _) in enumerate(ALL_METRICS, 3):
            med = pd.to_numeric(sub[key], errors="coerce").median()
            c = ws.cell(row=s, column=j, value=None if pd.isna(med) else ("Debt Free" if np.isinf(med) else float(med)))
            c.number_format = _FMT.get(key, "0.0")
        for j in range(1, 3 + 2 * n_metrics):
            c = ws.cell(row=s, column=j)
            c.fill = GREY if c.fill.fgColor.rgb in (None, "00000000") else c.fill
            c.font = Font(bold=True)
            c.border = Border(top=Side(style="medium", color="0B2239"))
        for j in range(3 + n_metrics, 3 + 2 * n_metrics):
            ws.cell(row=s, column=j).fill = GREY

        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 38
        for j in range(3, 3 + 2 * n_metrics):
            ws.column_dimensions[get_column_letter(j)].width = 12
        ws.freeze_panes = ws.cell(row=hdr + 1, column=3)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
