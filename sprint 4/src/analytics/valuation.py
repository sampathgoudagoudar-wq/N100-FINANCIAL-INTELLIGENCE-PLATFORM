"""Valuation module: FCF yield, sector-median P/E flags, 5yr median P/E (market_cap.xlsx)."""
from __future__ import annotations

import numpy as np
import pandas as pd

CAUTION_MULT, DISCOUNT_MULT = 1.5, 0.7


def compute_valuation(t: dict[str, pd.DataFrame], snapshot: pd.DataFrame) -> pd.DataFrame:
    """One row per company (latest market year).

    FCF_yield_pct = FCF / market_cap_crore x 100
    flag: P/E > sector median x 1.5 -> Caution ; P/E < sector median x 0.7 -> Discount ; else Fair
    (P/E missing -> 'N/A').
    """
    mc = t["market_cap"]
    latest = mc[mc["year"] == mc["year"].max()].set_index("company_id")
    last5 = mc[mc["year"] > mc["year"].max() - 5]
    pe5 = last5.groupby("company_id")["pe_ratio"].median().rename("pe_5yr_median")

    snap = snapshot.set_index("company_id")
    out = pd.DataFrame(index=snap.index)
    out["company_name"] = snap["company_name"]
    out["sector"] = snap["broad_sector"]
    out["pe"] = latest["pe_ratio"]
    out["pb"] = latest["pb_ratio"]
    out["ev_ebitda"] = latest["ev_ebitda"]
    out["market_cap_crore"] = latest["market_cap_crore"]
    out["fcf_cr"] = snap["fcf"]
    out["fcf_yield_pct"] = (snap["fcf"] / latest["market_cap_crore"].where(latest["market_cap_crore"] > 0) * 100)
    out["pe_5yr_median"] = pe5
    out["sector_median_pe"] = out.groupby("sector")["pe"].transform("median")
    out["pe_vs_sector_median_pct"] = (out["pe"] / out["sector_median_pe"] - 1) * 100
    flag = np.where(out["pe"] > out["sector_median_pe"] * CAUTION_MULT, "Caution",
            np.where(out["pe"] < out["sector_median_pe"] * DISCOUNT_MULT, "Discount", "Fair"))
    out["flag"] = np.where(out["pe"].isna(), "N/A", flag)
    return out.reset_index()


SUMMARY_COLUMNS = {
    "company_id": "company_id", "company_name": "company_name", "sector": "sector", "pe": "P/E", "pb": "P/B",
    "ev_ebitda": "EV/EBITDA", "fcf_yield_pct": "FCF_yield_pct", "pe_5yr_median": "5yr_median_PE",
    "pe_vs_sector_median_pct": "PE_vs_sector_median_pct", "flag": "flag",
}


def write_valuation_outputs(val: pd.DataFrame, xlsx_path, csv_path) -> None:
    from openpyxl.styles import Font, PatternFill
    summary = val[list(SUMMARY_COLUMNS)].rename(columns=SUMMARY_COLUMNS).round(2)
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Valuation", index=False)
        ws = xw.sheets["Valuation"]
        fills = {"Caution": "FFC7CE", "Discount": "C6EFCE", "Fair": "FFEB9C"}
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", start_color="0B2239", end_color="0B2239")
        col = list(summary.columns).index("flag") + 1
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=col)
            if cell.value in fills:
                cell.fill = PatternFill("solid", start_color=fills[cell.value], end_color=fills[cell.value])
        for i, name in enumerate(summary.columns, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 36 if name == "company_name" else 18
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
    flagged = val[val["flag"].isin(["Caution", "Discount"])]
    cols = ["company_id", "company_name", "sector", "pe", "sector_median_pe", "pe_vs_sector_median_pct",
            "pe_5yr_median", "pb", "ev_ebitda", "fcf_yield_pct", "market_cap_crore", "flag"]
    flagged[cols].round(2).to_csv(csv_path, index=False)
